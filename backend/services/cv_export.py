import io
from typing import Literal
import re

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer


def _is_heading(line: str) -> bool:
    if not line:
        return False
    cleaned = line.strip()
    if len(cleaned) > 45:
        return False
    if ":" in cleaned and len(cleaned.split()) > 6:
        return False
    return cleaned.upper() == cleaned


def _normalize_lines(cv_text: str) -> list[str]:
    return [line.rstrip() for line in cv_text.splitlines()]


def _extract_candidate_name(cv_text: str) -> str:
    for raw in _normalize_lines(cv_text):
        line = raw.strip()
        if not line:
            continue
        if _is_heading(line):
            continue
        return line
    return "ATS Optimized CV"


def build_export_filename(cv_text: str, export_format: Literal["pdf", "docx"]) -> str:
    name = _extract_candidate_name(cv_text)
    cleaned = re.sub(r"[^A-Za-z0-9 ]+", "", name).strip()
    cleaned = re.sub(r"\s+", "_", cleaned)
    if not cleaned:
        cleaned = "ATS_Optimized_CV"
    return f"{cleaned}_ATS_CV.{export_format}"


def export_cv_docx(cv_text: str) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Pt(48)
    section.bottom_margin = Pt(48)
    section.left_margin = Pt(52)
    section.right_margin = Pt(52)

    lines = _normalize_lines(cv_text)
    title_written = False

    for raw in lines:
        line = raw.strip()
        if not line:
            doc.add_paragraph("")
            continue

        if not title_written:
            p = doc.add_paragraph(line)
            p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
            run = p.runs[0]
            run.bold = True
            run.font.size = Pt(17)
            title_written = True
            continue

        if _is_heading(line):
            p = doc.add_paragraph(line)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.runs[0]
            run.bold = True
            run.font.size = Pt(12)
            continue

        if line.startswith(("- ", "* ")):
            p = doc.add_paragraph(line[2:].strip(), style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            if p.runs:
                p.runs[0].font.size = Pt(10.5)
            continue

        p = doc.add_paragraph(line)
        p.paragraph_format.space_after = Pt(3)
        if p.runs:
            p.runs[0].font.size = Pt(10.5)

    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def export_cv_pdf(cv_text: str) -> bytes:
    output = io.BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=A4,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        title="ATS Optimized CV",
    )

    styles = getSampleStyleSheet()
    normal = ParagraphStyle(
        "CVNormal",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10.8,
        leading=14.6,
        textColor=colors.HexColor("#0f172a"),
        spaceAfter=4,
    )
    heading = ParagraphStyle(
        "CVHeading",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=11.8,
        leading=15.8,
        textColor=colors.HexColor("#0b2a4a"),
        spaceBefore=11,
        spaceAfter=5,
    )
    title = ParagraphStyle(
        "CVTitle",
        parent=styles["Title"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22.5,
        alignment=1,
        textColor=colors.HexColor("#07111f"),
        spaceAfter=6,
    )
    subtitle = ParagraphStyle(
        "CVSubtitle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9.2,
        leading=12,
        alignment=1,
        textColor=colors.HexColor("#475569"),
        spaceAfter=8,
    )

    story = []
    lines = _normalize_lines(cv_text)
    title_written = False

    for raw in lines:
        line = raw.strip()
        if not line:
            story.append(Spacer(1, 5))
            continue

        safe_line = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

        if not title_written:
            story.append(Paragraph(safe_line, title))
            story.append(Paragraph("ATS-Optimized Resume", subtitle))
            story.append(
                HRFlowable(
                    width="100%",
                    thickness=0.8,
                    lineCap="round",
                    color=colors.HexColor("#94a3b8"),
                    spaceBefore=2,
                    spaceAfter=8,
                )
            )
            title_written = True
            continue

        if _is_heading(line):
            story.append(Paragraph(safe_line, heading))
            continue

        if line.startswith(("- ", "* ")):
            bullet = safe_line[2:].strip()
            story.append(Paragraph(f"• {bullet}", normal))
            continue

        story.append(Paragraph(safe_line, normal))

    doc.build(story)
    return output.getvalue()


def export_cv(cv_text: str, export_format: Literal["pdf", "docx"]) -> bytes:
    if export_format == "pdf":
        return export_cv_pdf(cv_text)
    if export_format == "docx":
        return export_cv_docx(cv_text)
    raise ValueError("Unsupported format. Use 'pdf' or 'docx'.")
